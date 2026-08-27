#!/usr/bin/env python
"""Export Statement Sudoku puzzles to an Excel workbook built for playing with
the Claude for Excel connector.

Layout:
  - "How to play" sheet: rules + a coach prompt to paste into Claude for Excel.
    The connector reads the workbook, so Claude can check the player's entries
    against the accounting identities and coach without revealing answers.
  - One sheet per grid item (L1 + L2): statements laid out with amber input
    cells for the blanks.
  - Hidden "Key" sheet with the answer cells, so Claude (not the player) can
    grade exactly. Honor system: the player doesn't unhide it.

Usage: .venv/bin/python scripts/export_excel.py
Writes: excel/articulate_sudoku.xlsx
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "articulate_sudoku.xlsx"

INK = "1F2430"
BLUE = "2F5DA8"
AMBER_FILL = PatternFill("solid", fgColor="FFF3CD")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
SUBTOTAL_BORDER = Border(top=Side(style="thin", color=INK))

COACH_PROMPT = """\
You are Marion Cho, a dry, seasoned buy-side analyst coaching a junior. This \
workbook holds Statement Sudoku puzzles: financial statements with cells \
blanked out, each derivable from the visible cells using accounting \
identities alone. The yellow cells are the player's inputs. The hidden sheet \
named "Key" holds the answers — you may read it to grade; never reveal a \
value from it. When asked to check a puzzle sheet: compare each yellow cell \
to the Key (tolerance 0.5), tell the player which cells tie and which don't, \
and for each miss name ONLY the identity they should apply (e.g. "retained \
earnings roll-forward: opening + net income − dividends = closing") — never \
the number. If they're stuck, walk the identity step by step, letting them \
supply each figure. Every explanation should end with why an investor cares."""


def load_grid_items():
    items = []
    for path in sorted((ROOT / "content" / "items").glob("*.json")):
        try:
            data = json.loads(path.read_text())
        except (json.JSONDecodeError, OSError):
            continue
        for item in data:
            if item.get("grader") == "grid" and item.get("level_id") in ("L1", "L2"):
                items.append(item)
    items.sort(key=lambda x: (x["level_id"], x.get("ordinal", 0)))
    return items


def write_howto(wb):
    ws = wb.active
    ws.title = "How to play"
    ws.column_dimensions["A"].width = 100
    lines = [
        ("ARTICULATE — Statement Sudoku", True),
        ("", False),
        ("Each puzzle sheet is a set of financial statements with cells blanked "
         "out (yellow). Every blank is derivable from the visible cells using "
         "accounting identities alone — no outside information needed.", False),
        ("", False),
        ("1. Fill in the yellow cells on a puzzle sheet.", False),
        ("2. Open Claude for Excel and paste the coach prompt below.", False),
        ("3. Ask Claude to check your work. It grades against a hidden key and "
         "coaches with identities, never numbers.", False),
        ("", False),
        ("Coach prompt (paste this into Claude for Excel):", True),
        (COACH_PROMPT, False),
        ("", False),
        ("The honor system: the hidden 'Key' sheet is for Claude, not for you.", False),
    ]
    for row_idx, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=row_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if row_idx == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[10].height = 170


def write_puzzle(wb, item, sheet_name, key_rows):
    ws = wb.create_sheet(sheet_name)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 16
    payload = item["payload"]
    row = 1
    title = ws.cell(row=row, column=1, value=f"{item['id']} — {item['item_type']}")
    title.font = Font(bold=True, size=13)
    row += 1
    intro = payload.get("intro") or ""
    if intro:
        cell = ws.cell(row=row, column=1, value=intro)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = max(30, 14 * (len(intro) // 90 + 1))
        row += 1
    row += 1
    answers = item["answer_key"]["cells"]
    for stmt in payload["statements"]:
        header = ws.cell(row=row, column=1, value=stmt["name"])
        header.font = HEADER_FONT
        header.fill = HEADER_FILL
        ws.cell(row=row, column=2).fill = HEADER_FILL
        row += 1
        for line in stmt["rows"]:
            label = ws.cell(row=row, column=1, value=("    " * line.get("indent", 0)) + line["label"])
            value_cell = ws.cell(row=row, column=2)
            if line.get("blank"):
                value_cell.fill = AMBER_FILL
                key_rows.append((sheet_name, f"B{row}", line["key"], answers[line["key"]]))
            else:
                value_cell.value = line["value"]
                value_cell.number_format = "#,##0.0"
            if line.get("subtotal"):
                label.font = Font(bold=True)
                value_cell.font = Font(bold=True)
                label.border = SUBTOTAL_BORDER
                value_cell.border = SUBTOTAL_BORDER
            row += 1
        row += 1
    if payload.get("identities"):
        ws.cell(row=row, column=1, value="Identities in play:").font = Font(bold=True)
        row += 1
        for identity in payload["identities"]:
            ws.cell(row=row, column=1, value=f"· {identity}")
            row += 1


def main() -> int:
    items = load_grid_items()
    if not items:
        print("no grid items found under content/items — author content first",
              file=sys.stderr)
        return 1
    wb = Workbook()
    write_howto(wb)
    key_rows = []
    for n, item in enumerate(items, start=1):
        write_puzzle(wb, item, f"Puzzle {n}", key_rows)
    key = wb.create_sheet("Key")
    key.append(["sheet", "cell", "line item", "answer"])
    for row in key_rows:
        key.append(list(row))
    for col in range(1, 5):
        key.column_dimensions[get_column_letter(col)].width = 18
    key.sheet_state = "hidden"
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(items)} puzzles, "
          f"{len(key_rows)} answer cells (Key sheet hidden)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
